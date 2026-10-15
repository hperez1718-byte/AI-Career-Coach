"""🎓 AI Career Coach
A free, open-source AI job-search assistant. Run it locally with your Google
Gemini API key — no app account or saved session history.

HOW IT FLOWS:
  Sidebar ......... Paste your Gemini API key, then load the resume ONCE
                    (upload a PDF/Word/txt or paste). Every tab uses it for
                    this browser session.
  Tab 1 ATS Match . Paste a job posting → runs against the sidebar resume.
                    Verdict: match or not, why / why not, how to close the gap.
  Tab 2 Optimizer . Takes the job + the ATS advice and builds a tailored
                    resume + cover letter from the sidebar resume.
  Tab 3 Mock ..... Paste (or receive) a job posting → company briefing,
                    personal advice, and 5 STAR interview Q&As built from
                    the sidebar resume.

Setup:
  pip install -r requirements.txt
  streamlit run app.py
  Paste your free Gemini API key in the sidebar (https://aistudio.google.com/apikey)
"""

import streamlit as st
import os
import html
from google import genai
from docx import Document
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from dotenv import load_dotenv

# --- API key (.env is the fallback; sidebar input takes priority) ----------
_script_dir = os.path.dirname(os.path.abspath(__file__))
_env_path = os.path.join(_script_dir, ".env")
if os.path.isfile(_env_path):
    load_dotenv(_env_path)
ENV_KEY = os.getenv("GOOGLE_API_KEY", "")

# A lighter model is a better default for a classroom: lower cost and intended
# for high-volume work. The higher-quality model remains available in the UI.
MODEL = "gemini-3.5-flash-lite"
FALLBACK_MODEL = "gemini-3.8-flash"
MODEL_LABELS = {
    MODEL: "Classroom / fast — Gemini 3.5 Flash-Lite",
    FALLBACK_MODEL: "Higher quality — Gemini 3.8 Flash",
}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_RESUME_CHARS = 30_000
MAX_JOB_DESCRIPTION_CHARS = 25_000

st.set_page_config(page_title="AI Career Coach", page_icon="🎓", layout="wide")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Orbitron:wght@600;800&display=swap');
header[data-testid="stHeader"] { visibility: hidden; }
.stApp { background: linear-gradient(180deg, #eff6ff 0%, #ffffff 55%); color: #0f172a; }
h1 { font-family: 'Orbitron', sans-serif !important; font-weight: 800 !important;
     background: linear-gradient(90deg, #0c4a6e, #0284c7); -webkit-background-clip: text;
     -webkit-text-fill-color: transparent; background-clip: text; letter-spacing: 3px; }
h2, h3 { color: #0c4a6e !important; font-weight: 700; }
section[data-testid="stSidebar"] { background: linear-gradient(180deg, #e0f2fe 0%, #f0f9ff 100%); border-right: 2px solid #bae6fd; }
section[data-testid="stSidebar"] [data-testid="stCaptionContainer"],
section[data-testid="stSidebar"] [data-testid="stCaptionContainer"] p,
section[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p,
section[data-testid="stSidebar"] label { color: #334155 !important; }
[data-testid="stTextArea"] textarea, [data-testid="stTextInput"] input {
    background-color: #ffffff !important; color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    border: 1px solid #7dd3fc !important; border-radius: 10px !important;
    box-shadow: 0 1px 3px rgba(2,132,199,0.10);
}
[data-testid="stFileUploader"] section {
    background: #ffffff !important; border: 1px dashed #38bdf8 !important;
    border-radius: 10px !important;
}
[data-testid="stFileUploader"] section small,
[data-testid="stFileUploader"] section span { color: #334155 !important; }
[data-testid="stFileUploader"] button {
    background: #e0f2fe !important; color: #0c4a6e !important;
    border: 1px solid #7dd3fc !important; box-shadow: none !important;
}
[data-testid="stAlert"] {
    background: #fef3c7 !important; border: 1px solid #fcd34d !important;
    color: #92400e !important;
}
[data-testid="stAlert"] p, [data-testid="stAlert"] div,
[data-testid="stAlert"] span { color: #92400e !important; }
[data-testid="stAlert"] a { color: #0c4a6e !important; font-weight: 700; }
.class-card { background: #ffffff; border: 1px solid #bae6fd; border-radius: 16px; padding: 28px;
              box-shadow: 0 4px 20px rgba(2,132,199,0.08); margin-bottom: 20px; }
div.stButton > button, div.stDownloadButton > button {
    background: linear-gradient(90deg, #0369a1, #38bdf8) !important;
    color: #fff !important; font-weight: 700; border: none !important;
    border-radius: 10px; width: 100%; letter-spacing: 0.5px;
    box-shadow: 0 2px 8px rgba(2,132,199,0.30);
}
div.stButton > button:hover, div.stDownloadButton > button:hover {
    filter: brightness(1.1); box-shadow: 0 4px 14px rgba(2,132,199,0.45);
}
.pipe { text-align: center; color: #0284c7; font-weight: 600; margin: 4px 0 16px 0; letter-spacing: 1px; }
.accent-line { height: 3px; background: linear-gradient(90deg, #ef4444, #f97316, #38bdf8);
               border-radius: 2px; margin: 6px 0 14px 0; }
.brand-footer { text-align: center; padding: 16px; color: #64748b; font-size: 0.85em; }
.brand-footer b { color: #0c4a6e; }
</style>
""", unsafe_allow_html=True)

# --- Session state ---------------------------------------------------------
if "active_module" not in st.session_state:
    st.session_state.active_module = "ats"
for key in ("user_api_key", "my_resume", "sb_last_file",
            "ats_jd_text", "ats_output", "ats_risk_output", "ats_keyword_report",
            "opt_jd_text", "opt_ats_advice", "opt_resume_output", "opt_cover_output", "opt_summary",
            "mock_jd_text", "mock_summary", "mock_output"):
    if key not in st.session_state:
        st.session_state[key] = ""
for _pk, _dv in (("practice_questions", []), ("practice_feedback", []),
                 ("practice_idx", 0), ("practice_active", False),
                 ("practice_summary", "")):
    if _pk not in st.session_state:
        st.session_state[_pk] = _dv
if "selected_model" not in st.session_state:
    st.session_state.selected_model = MODEL

# --- Helpers ---------------------------------------------------------------
def clear_generated_state():
    """Prevent results for an old resume or job posting from being reused."""
    for key in ("ats_output", "ats_risk_output", "ats_keyword_report",
                "opt_resume_output", "opt_cover_output", "opt_summary",
                "mock_summary", "mock_output", "practice_summary"):
        st.session_state[key] = ""
    st.session_state.practice_questions = []
    st.session_state.practice_feedback = []
    st.session_state.practice_idx = 0
    st.session_state.practice_active = False

def set_resume(resume_text, file_id=None):
    """Store one resume for the active browser session and reset old work."""
    if resume_text != st.session_state.my_resume:
        st.session_state.my_resume = resume_text
        clear_generated_state()
    if file_id is not None:
        st.session_state.sb_last_file = file_id

def ai_call(prompt, model=None):
    def _gen(m):
        response = client.models.generate_content(model=m, contents=prompt)
        if not response.text:
            raise RuntimeError("Gemini returned an empty response.")
        return response.text

    preferred = model or st.session_state.selected_model
    models = []
    for candidate in (preferred, MODEL, FALLBACK_MODEL):
        if candidate not in models:
            models.append(candidate)
    quota_error = None
    last_error = None
    for m in models:
        try:
            return _gen(m)
        except Exception as e:
            msg = str(e)
            last_error = e
            if "429" in msg or "RESOURCE_EXHAUSTED" in msg or "quota" in msg.lower():
                quota_error = msg
                continue
            if "503" in msg or "UNAVAILABLE" in msg or "overloaded" in msg.lower():
                continue
            if "404" in msg or "NOT_FOUND" in msg:
                continue
            raise
    if quota_error:
        safe_error = quota_error.replace(API_KEY, "[redacted]") if API_KEY else quota_error
        st.error("⚠️ Gemini rejected this request because this Google project has reached, or currently has no access to, its API quota.")
        st.markdown("[Check this project's active Gemini rate limits in AI Studio](https://aistudio.google.com/rate-limit)")
        st.caption("Changing to a different API key from the same Google project will not add capacity. Try the Classroom / fast model, wait for the project quota to reset, or use a project with available quota.")
        with st.expander("Technical details", expanded=False):
            st.code(safe_error)
        st.stop()
    if last_error:
        raise last_error
    raise RuntimeError("Gemini did not return a response.")

def make_pdf(text_content):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    bs = ParagraphStyle('BD', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13, textColor=colors.HexColor('#1e293b'), spaceAfter=6)
    hs = ParagraphStyle('HS', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=colors.HexColor('#0369a1'), spaceBefore=10, spaceAfter=4)
    story = []
    for line in text_content.split(chr(10)):
        clean = line.strip()
        if not clean:
            continue
        fmt = html.escape(clean.replace('**', '').replace('###', '').strip())
        if fmt.isupper() and len(fmt) < 50:
            story.append(Paragraph(fmt, hs))
        elif fmt[:1] in ('*', '-', '•'):
            story.append(Paragraph('• ' + fmt.lstrip('*-• ').strip(), bs))
        else:
            story.append(Paragraph(fmt, bs))
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def make_docx(text_content):
    doc = Document()
    for line in text_content.split(chr(10)):
        clean = line.strip()
        if not clean:
            continue
        fmt = clean.replace('**', '').replace('###', '').strip()
        if fmt.isupper() and len(fmt) < 50:
            doc.add_heading(fmt, level=2)
        elif fmt[:1] in ('*', '-', '•'):
            doc.add_paragraph(fmt.lstrip('*-• ').strip(), style='List Bullet')
        else:
            doc.add_paragraph(fmt)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def extract_text_from_upload(uploaded_file):
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        st.error("Please upload a resume smaller than 10 MB.")
        return ""
    name = uploaded_file.name.lower()
    try:
        if name.endswith(".txt"):
            text = uploaded_file.read().decode("utf-8", errors="ignore")
        elif name.endswith(".docx"):
            doc = Document(uploaded_file)
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(uploaded_file)
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        else:
            st.warning("Please upload a .pdf, .docx, or .txt file.")
            return ""
    except Exception:
        st.error("This file could not be read. Try exporting the resume as a text-based PDF, DOCX, or TXT file.")
        return ""
    if len(text) > MAX_RESUME_CHARS:
        st.error("This resume contains too much text. Please upload a shorter version (about 30,000 characters or less).")
        return ""
    if not text.strip():
        st.error("No readable text was found. If this is a scanned PDF, upload a DOCX or TXT version instead.")
        return ""
    return text

def extract_keywords(jd_text, resume_text):
    ep = 'Extract the top 15-20 most important skills, tools, technologies, certifications, and experience requirements from this Job Description. List each as a short phrase. One per line.\n\nJob Description:\n' + jd_text
    kw = [l.strip() for l in ai_call(ep).split(chr(10)) if l.strip()]
    rl = resume_text.lower()
    return kw, [k for k in kw if k.lower() in rl], [k for k in kw if k.lower() not in rl]

def get_resume():
    r = st.session_state.my_resume.strip()
    if not r:
        st.warning("📝 No resume loaded yet — upload or paste one in the **sidebar** first.")
        return None
    if len(r) > MAX_RESUME_CHARS:
        st.warning("📝 Please shorten the resume to 30,000 characters or less before running an AI feature.")
        return None
    return r

# --- Sidebar: API key + resume ---------------------------------------------
with st.sidebar:
    st.markdown("## 🎓 AI Career Coach")
    st.caption("Load your key, load one resume, then work the pipeline.")
    st.markdown("### 🔑 API Key")
    st.caption("Paste your free Gemini key, or leave blank to use a `.env` file.")
    _ki = st.text_input("Gemini API key", type="password", value=st.session_state.user_api_key,
                        key="key_input", label_visibility="collapsed", placeholder="Paste key here...")
    k1, k2 = st.columns(2)
    with k1:
        if st.button("Save", key="key_save"):
            st.session_state.user_api_key = _ki.strip()
            st.rerun()
    with k2:
        if st.button("Clear", key="key_clear"):
            st.session_state.user_api_key = ""
            st.rerun()
    if (st.session_state.user_api_key.strip() or ENV_KEY):
        st.success("API Key Loaded ✓")
    else:
        st.warning("No API key yet")
    st.markdown("### 🧠 AI Model")
    st.selectbox("Choose an AI model", options=list(MODEL_LABELS),
                 format_func=lambda value: MODEL_LABELS[value], key="selected_model",
                 label_visibility="collapsed")
    st.caption("Classroom / fast is the recommended default. If a project reaches its Google quota, changing models may not restore access.")
    st.markdown("---")
    st.markdown("### 📄 My Resume")
    st.caption("Load it once here — every tab uses this resume.")
    up = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"], key="sb_upload", label_visibility="collapsed")
    if up is not None:
        fid = up.name + "_" + str(up.size)
        if st.session_state.get("sb_last_file") != fid:
            txt = extract_text_from_upload(up)
            if txt.strip():
                set_resume(txt, fid)
                st.rerun()
    if st.button("🗑️ Clear resume", key="sb_clear"):
        set_resume("")
        st.session_state.sb_last_file = ""
        st.rerun()
    st.text_area("Resume text", key="my_resume", on_change=clear_generated_state,
        height=220, placeholder="Or paste the resume here as plain text...", label_visibility="collapsed")
    if st.session_state.my_resume.strip():
        st.success("Resume ready (" + str(len(st.session_state.my_resume.split())) + " words) ✓")
        st.caption("Used by every tab in this browser session. New resume = fresh results.")
    else:
        st.warning("No resume loaded yet")
    with st.expander("🏫 Classroom workflow", expanded=False):
        st.markdown("""1. Load a resume once.  
2. Run **ATS Match** with one job posting.  
3. Send the job to **Resume Optimizer**.  
4. Send it to **Mock Interview** and practice.""")
        st.caption("Resume and results stay in this browser session only. Download work before closing or refreshing the tab.")
    st.caption("Privacy: your resume and job posting are sent to Google's Gemini API when you run AI features. This app does not write new session logs.")

# --- Resolve key + create client -------------------------------------------
API_KEY = st.session_state.user_api_key.strip() or ENV_KEY
if not API_KEY:
    st.warning("🔑 **Add your Gemini API key in the sidebar to get started.**\n\nIt's free — [get one here](https://aistudio.google.com/apikey), paste it above, and hit **Save**.")
    st.stop()
client = genai.Client(api_key=API_KEY)

# --- Header + nav ----------------------------------------------------------
st.title("🎓 AI Career Coach")
st.markdown('<div class="accent-line"></div>', unsafe_allow_html=True)
st.markdown("Load the resume in the sidebar, then work the pipeline: match it → tailor it → practice it.")
st.markdown('<div class="pipe">📄 Sidebar: resume &nbsp;→&nbsp; 📈 Tab 1: match &nbsp;→&nbsp; 🔄 Tab 2: tailor &nbsp;→&nbsp; 💡 Tab 3: interview</div>', unsafe_allow_html=True)

nc1, nc2, nc3 = st.columns(3)
with nc1:
    if st.button("📈 ATS Match", key="nav_ats"):
        st.session_state.active_module = "ats"
        st.rerun()
with nc2:
    if st.button("🔄 Resume Optimizer", key="nav_opt"):
        st.session_state.active_module = "optimizer"
        st.rerun()
with nc3:
    if st.button("💡 Mock Interview", key="nav_mock"):
        st.session_state.active_module = "mock"
        st.rerun()
st.markdown("---")

# --- Tab 1: ATS Match -------------------------------------------------------
if st.session_state.active_module == "ats":
    st.markdown('<div class="class-card">', unsafe_allow_html=True)
    st.subheader("📈 ATS Match Simulator")
    st.markdown("Paste a job posting — it runs against the resume in the sidebar: **match or not, why / why not**, and how to close the gap.")
    resume = get_resume()
    if resume:
        st.session_state.ats_jd_text = st.text_area("Job posting", value=st.session_state.ats_jd_text,
            height=180, placeholder="Paste the full job description here...", label_visibility="collapsed")
        if st.button("📊 Run Match"):
            jd = st.session_state.ats_jd_text.strip()
            if not jd:
                st.warning("Please paste a job posting.")
            elif len(jd) > MAX_JOB_DESCRIPTION_CHARS:
                st.warning("Please shorten the job posting to 25,000 characters or less.")
            else:
                st.session_state.ats_output = ""
                st.session_state.ats_risk_output = ""
                st.session_state.ats_keyword_report = ""
                with st.spinner("Extracting key requirements..."):
                    all_kw, matched, missing = extract_keywords(jd, resume)
                st.session_state.ats_keyword_report = ("**Matched (" + str(len(matched)) + "):** "
                    + (", ".join(matched[:20]) if matched else "none")
                    + "\n\n**Missing (" + str(len(missing)) + "):** "
                    + (", ".join(missing[:20]) if missing else "None — all found!"))
                with st.spinner("Scoring the match..."):
                    combined = ai_call(
                        "You are an ATS simulator. Decide: is this candidate a match for this job? Be honest.\n\n"
                        + "Keywords matched: " + ", ".join(matched)
                        + "\nKeywords missing: " + ", ".join(missing)
                        + "\n\nJob Description:\n" + jd + "\n\nResume:\n" + resume
                        + "\n\nPart 1 — give: 1) **Verdict** — Match or Not a Match (one line), 2) **Match Score (%)**, "
                        + "3) **Why** — what aligns (2-3 sentences), 4) **Why not** — the most important gaps, in priority order."
                        + "\n\nThen write exactly this separator on its own line: ===COACHING_TIPS==="
                        + "\n\nPart 2 — coaching tips for the student: 4-5 concrete steps to become a stronger match "
                        + "(exact keywords to add, experience to reframe, skills to learn). Be encouraging but honest.")
                    if "===COACHING_TIPS===" in combined:
                        parts = combined.split("===COACHING_TIPS===", 1)
                        st.session_state.ats_output = parts[0].strip()
                        st.session_state.ats_risk_output = parts[1].strip()
                    else:
                        st.session_state.ats_output = combined
                        st.session_state.ats_risk_output = ""
        if st.session_state.ats_keyword_report:
            st.markdown("**🔍 Keyword Report**")
            st.markdown(st.session_state.ats_keyword_report)
        if st.session_state.ats_output:
            st.markdown("---")
            st.subheader("📊 Match Result")
            st.markdown(st.session_state.ats_output)
            if st.session_state.ats_risk_output:
                st.markdown("**💡 How to Become a Match**")
                st.markdown(st.session_state.ats_risk_output)
            st.markdown("---")
            s1, s2 = st.columns(2)
            with s1:
                if st.button("➡️ Send to Resume Optimizer"):
                    st.session_state.opt_jd_text = st.session_state.ats_jd_text
                    st.session_state.opt_ats_advice = st.session_state.ats_risk_output or st.session_state.ats_output
                    st.session_state.active_module = "optimizer"
                    st.rerun()
            with s2:
                if st.button("💡 Send job to Mock Interview", key="ats_to_mock"):
                    st.session_state.mock_jd_text = st.session_state.ats_jd_text
                    st.session_state.active_module = "mock"
                    st.rerun()
            d1, d2 = st.columns(2)
            with d1:
                st.download_button("📥 Download PDF", data=make_pdf(st.session_state.ats_output + "\n\n" + st.session_state.ats_risk_output), file_name="ats_report.pdf")
            with d2:
                st.download_button("📝 Download Word", data=make_docx(st.session_state.ats_output + "\n\n" + st.session_state.ats_risk_output), file_name="ats_report.docx")
    st.markdown('</div>', unsafe_allow_html=True)

# --- Tab 2: Resume Optimizer ------------------------------------------------
elif st.session_state.active_module == "optimizer":
    st.markdown('<div class="class-card">', unsafe_allow_html=True)
    st.subheader("🔄 Resume Optimizer")
    st.markdown("Builds a tailored resume + cover letter from the sidebar resume. If you sent ATS advice over, it's applied automatically.")
    resume = get_resume()
    if resume:
        if st.session_state.opt_ats_advice:
            with st.expander("📋 ATS notes applied to this run", expanded=False):
                st.markdown(st.session_state.opt_ats_advice)
        st.session_state.opt_jd_text = st.text_area("Job posting", value=st.session_state.opt_jd_text,
            height=180, placeholder="Paste the full job description here...", label_visibility="collapsed")
        if st.button("✨ Generate Tailored Resume"):
            jd = st.session_state.opt_jd_text.strip()
            if not jd:
                st.warning("Please paste a job posting.")
            elif len(jd) > MAX_JOB_DESCRIPTION_CHARS:
                st.warning("Please shorten the job posting to 25,000 characters or less.")
            else:
                st.session_state.opt_summary = ""
                st.session_state.opt_resume_output = ""
                st.session_state.opt_cover_output = ""
                advice = st.session_state.opt_ats_advice.strip()
                advice_block = ("Apply this feedback from an ATS scan of this resume:\n" + advice + "\n\n") if advice else ""
                with st.spinner("Analyzing the role..."):
                    st.session_state.opt_summary = ai_call(
                        "Summarize this job: company, what they care about most, top 3-4 requirements.\n\n" + jd)
                with st.spinner("Writing tailored resume..."):
                    st.session_state.opt_resume_output = ai_call(
                        "You are an expert resume writer. Tailor this resume to the job description.\n\n"
                        + advice_block
                        + "Rules:\n- Reframe and reorder the candidate's REAL experience to match the JD's keywords and priorities.\n"
                        + "- NEVER invent jobs, skills, certifications, or achievements not in the original resume.\n"
                        + "- Keep the candidate's contact header exactly as written.\n"
                        + "- Output ONLY the resume. No cover letter, no commentary.\n\nJob Description:\n" + jd
                        + "\n\nOriginal Resume:\n" + resume)
                with st.spinner("Writing cover letter..."):
                    st.session_state.opt_cover_output = ai_call(
                        "Write a professional cover letter (under 300 words) for this candidate and job. "
                        + "Draw ONLY on the real resume — do not invent experience. Output ONLY the letter.\n\nJob Description:\n" + jd
                        + "\n\nResume:\n" + resume)
        if st.session_state.opt_summary:
            st.markdown("**🏢 Role Breakdown**")
            st.markdown(st.session_state.opt_summary)
        if st.session_state.opt_resume_output:
            st.markdown("---")
            st.subheader("📄 Tailored Resume")
            st.markdown(st.session_state.opt_resume_output)
            c1, c2 = st.columns(2)
            with c1:
                st.download_button("📥 Resume PDF", data=make_pdf(st.session_state.opt_resume_output), file_name="tailored_resume.pdf")
            with c2:
                st.download_button("📝 Resume Word", data=make_docx(st.session_state.opt_resume_output), file_name="tailored_resume.docx")
        if st.session_state.opt_cover_output:
            st.markdown("---")
            st.subheader("✏️ Cover Letter")
            st.markdown(st.session_state.opt_cover_output)
            c1, c2 = st.columns(2)
            with c1:
                st.download_button("📥 Letter PDF", data=make_pdf(st.session_state.opt_cover_output), file_name="cover_letter.pdf")
            with c2:
                st.download_button("📝 Letter Word", data=make_docx(st.session_state.opt_cover_output), file_name="cover_letter.docx")
        if st.session_state.opt_jd_text.strip():
            st.markdown("---")
            if st.button("💡 Send job to Mock Interview", key="opt_to_mock"):
                st.session_state.mock_jd_text = st.session_state.opt_jd_text
                st.session_state.active_module = "mock"
                st.rerun()
    st.markdown('</div>', unsafe_allow_html=True)

# --- Tab 3: Mock Interview --------------------------------------------------
elif st.session_state.active_module == "mock":
    st.markdown('<div class="class-card">', unsafe_allow_html=True)
    st.subheader("💡 Mock Interview — STAR Prep")
    st.markdown("Company briefing, personal advice, and **5 interview questions** with STAR answers built from the sidebar resume.")
    resume = get_resume()
    if resume:
        st.session_state.mock_jd_text = st.text_area("Job posting", value=st.session_state.mock_jd_text,
            height=180, placeholder="Paste the full job description here — or send one over from another tab...", label_visibility="collapsed")
        if st.button("🚀 Generate Interview Prep"):
            jd = st.session_state.mock_jd_text.strip()
            if not jd:
                st.warning("Please paste a job posting.")
            elif len(jd) > MAX_JOB_DESCRIPTION_CHARS:
                st.warning("Please shorten the job posting to 25,000 characters or less.")
            else:
                st.session_state.mock_summary = ""
                st.session_state.mock_output = ""
                st.session_state.practice_questions = []
                st.session_state.practice_feedback = []
                st.session_state.practice_idx = 0
                st.session_state.practice_active = False
                st.session_state.practice_summary = ""
                with st.spinner("Analyzing the job posting..."):
                    st.session_state.mock_summary = ai_call(
                        "Analyze this job posting and write an interview briefing in two parts.\n\n"
                        + "PART 1 — COMPANY CONTEXT FROM THE JOB POSTING:\n"
                        + "- Company name, industry, and what it appears to do\n"
                        + "- Culture clues from the posting (pace, values, team style)\n"
                        + "- Top 2-3 things this hire must deliver\n\n"
                        + "PART 2 — PERSONAL ADVICE FOR THE CANDIDATE:\n"
                        + "- How they should present themselves for THIS company (2-3 sentences)\n"
                        + "- Which parts of their background to emphasize most\n"
                        + "- 3 practical confidence and presentation tips tailored to them\n\n"
                        + "Job Description:\n" + jd + "\n\nCandidate Resume:\n" + resume)
                with st.spinner("Building your STAR questions..."):
                    st.session_state.mock_output = ai_call(
                        "You are an expert interviewer and career coach. Using the job description and the candidate's resume below, generate 5 interview questions: "
                        + "2 technical/role-specific, 2 personal/behavioral (like 'Tell me about yourself' or 'Why do you want this job?'), and 1 curveball.\n\n"
                        + "For EACH question provide:\n"
                        + "- **Question:** the question itself\n"
                        + "- **Why they ask it:** what the interviewer is really evaluating (2-3 sentences)\n"
                        + "- **STAR sample answer:** a full Situation → Task → Action → Result narrative grounded ONLY in the candidate's real resume history. Do not invent experience.\n\n"
                        + "Job Description:\n" + jd + "\n\nCandidate Resume:\n" + resume)
        if st.session_state.mock_summary:
            st.markdown("---")
            st.subheader("🏢 Company Briefing & Personal Advice")
            st.markdown(st.session_state.mock_summary)
        if st.session_state.mock_output:
            st.markdown("---")
            st.subheader("❓ 5 STAR Questions & Answers")
            st.markdown(st.session_state.mock_output)
            _guide = ((st.session_state.mock_summary + "\n\n") if st.session_state.mock_summary else "") + st.session_state.mock_output
            st.download_button("📥 Download Interview Guide", data=make_pdf(_guide), file_name="interview_guide.pdf")
        st.markdown("---")
        st.subheader("🎤 Interactive Practice Mode")
        st.markdown("Don't just read sample answers — practice your own. The AI plays interviewer and coaches each response.")
        if not st.session_state.practice_active:
            if st.button("🎤 Start Practice Session", key="prac_start"):
                _jd_p = st.session_state.mock_jd_text.strip()
                if not _jd_p:
                    st.warning("Paste a job posting above first.")
                else:
                    import re as _re
                    with st.spinner("Preparing your questions..."):
                        _q_text = ai_call(
                            "Generate 5 interview questions for this job: 2 technical/role-specific, "
                            "2 personal/behavioral, 1 curveball. Output ONLY the 5 questions as a numbered list. "
                            "No answers, no explanations.\n\nJob Description:\n" + _jd_p
                            + "\n\nCandidate Resume:\n" + resume)
                    _qs = []
                    for _line in _q_text.split(chr(10)):
                        _s = _line.strip()
                        if not _s:
                            continue
                        _s = _re.sub(r"^\d+[\.\)]\s*", "", _s)
                        if _s:
                            _qs.append(_s)
                    st.session_state.practice_questions = _qs[:5]
                    st.session_state.practice_idx = 0
                    st.session_state.practice_feedback = []
                    st.session_state.practice_summary = ""
                    st.session_state.practice_active = True
                    st.rerun()
        if st.session_state.practice_active:
            _qs = st.session_state.practice_questions
            _idx = st.session_state.practice_idx
            _total = len(_qs) if _qs else 5
            st.progress(min(_idx / _total, 1.0))
            for _i in range(_idx):
                with st.expander(f"✅ Q{_i+1}: {_qs[_i][:70]}", expanded=False):
                    st.markdown("**Your answer:**")
                    st.markdown(st.session_state.get(f"prac_ans_{_i}", "_No answer recorded._"))
                    st.markdown("**Coach feedback:**")
                    st.markdown(st.session_state.practice_feedback[_i] if _i < len(st.session_state.practice_feedback) else "")
            if _qs and _idx < len(_qs):
                st.markdown(f"**Question {_idx+1} of {len(_qs)}**")
                st.markdown("#### " + _qs[_idx])
                _ans = st.text_area("Your answer — type it as you'd say it", key=f"prac_ans_{_idx}", height=150,
                                    placeholder="Say it in your head, then type it here...")
                if st.button("Submit Answer", key=f"prac_submit_{_idx}"):
                    if not _ans.strip():
                        st.warning("Type your answer first.")
                    else:
                        with st.spinner("Coaching your answer..."):
                            _fb = ai_call(
                                "You are an interview coach giving feedback on a practice answer.\n\n"
                                f"Question: {_qs[_idx]}\nCandidate's answer: {_ans}\n"
                                f"Candidate background: {resume}\n\n"
                                "Give:\n1. **Score:** X/10\n"
                                "2. **What worked:** 2 specific points\n"
                                "3. **To improve:** 2 specific, actionable points\n"
                                "4. **Try this:** one sentence showing a stronger angle\n\n"
                                "Be encouraging but honest. Under 200 words total.")
                        st.session_state.practice_feedback.append(_fb)
                        st.session_state.practice_idx = _idx + 1
                        st.rerun()
            elif _qs:
                st.success("🎉 Practice complete — nice work!")
                if not st.session_state.practice_summary:
                    with st.spinner("Writing your coaching summary..."):
                        _compiled = ""
                        for _i, _q in enumerate(_qs):
                            _compiled += f"Q{_i+1}: {_q}\nA: {st.session_state.get(f'prac_ans_{_i}', '')}\n"
                            _compiled += f"Feedback: {st.session_state.practice_feedback[_i] if _i < len(st.session_state.practice_feedback) else ''}\n\n"
                        st.session_state.practice_summary = ai_call(
                            "Based on this interview practice session, write a brief coaching summary.\n\n"
                            + _compiled
                            + "\nGive: 1) **Overall readiness** (1-2 sentences), 2) **Biggest strength**, "
                            + "3) **#1 thing to work on** before the real interview. Be direct and encouraging.")
                        st.rerun()
                else:
                    st.markdown("**📋 Your Coaching Summary**")
                    st.markdown(st.session_state.practice_summary)
                    st.download_button("📥 Download Practice Report",
                                        data=make_pdf(st.session_state.practice_summary),
                                        file_name="practice_report.pdf")
            if st.button("🔄 Restart Practice", key="prac_restart"):
                st.session_state.practice_active = False
                st.session_state.practice_idx = 0
                st.session_state.practice_questions = []
                st.session_state.practice_feedback = []
                st.session_state.practice_summary = ""
                st.rerun()
    st.markdown('</div>', unsafe_allow_html=True)

st.markdown("---")
st.markdown('<div class="brand-footer">🎓 <b>AI Career Coach</b><br>Built with 💙 by <b>hperez1718-byte</b> · Your text is sent to Google\'s Gemini API only when you run an AI feature. This app keeps no new session logs.</div>', unsafe_allow_html=True)
